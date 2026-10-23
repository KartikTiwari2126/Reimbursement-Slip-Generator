import streamlit as st
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from io import BytesIO
from datetime import date
from decimal import Decimal, InvalidOperation

# -------------------------------------------------------------------
# CONSTANT DETAILS — edit these once if committee details ever change.
# -------------------------------------------------------------------
SENDER_NAME = "Aadarsh Kumar"
SENDER_ID = "MS25A081"
SENDER_DEPARTMENT = "Department of Management Studies"
INSTITUTE = "Indian Institute of Technology Madras"
ADDRESS = "Chennai- 600036"
SIGNATURE_LINE = f"{SENDER_NAME} ({SENDER_ID})"
COMMITTEE = "Sports and Cultural Committee, DoMS"

RECIPIENT_LINES = [
    "The Head of the Department",
    "Department of Management Studies",
    "Indian Institute of Technology Madras",
    "Chennai- 600036",
]

# Bank details are included only in the Payment to Self letter.
BANK_NAME = "HDFC Bank"
BANK_ACCOUNT_NUMBER = "50100354506601"
BANK_ACCOUNT_HOLDER = "AADARSH KUMAR"
BANK_BRANCH = "KILPAUK"
BANK_IFSC = "HDFC0000124"


def money(value: Decimal) -> str:
    """Format INR amounts without unnecessary decimal places."""
    if value == value.quantize(Decimal("1")):
        return f"Rs. {value:,.0f}/-"
    return f"Rs. {value:,.2f}"


def parse_amount(value: str, field_name: str) -> Decimal:
    try:
        amount = Decimal(value.strip().replace(",", "").replace("₹", "").replace("Rs.", "").replace("Rs", "").strip())
        if amount < 0:
            raise ValueError
        return amount
    except (InvalidOperation, ValueError, AttributeError):
        raise ValueError(f"Enter a valid non-negative amount for {field_name}.")


def set_document_style(doc: Document):
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(6)


def add_line(doc: Document, text: str = "", bold: bool = False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    return p


def create_letter(payment_type, letter_date, event_name, purchase_date, items,
                  total_amount, approved_budget, vendor_name=""):
    doc = Document()
    set_document_style(doc)

    add_line(doc, "From:")
    add_line(doc, SENDER_NAME)
    add_line(doc, SENDER_DEPARTMENT)
    add_line(doc, INSTITUTE)
    add_line(doc, ADDRESS)
    add_line(doc, f"Date: {letter_date.strftime('%d-%m-%Y')}")
    doc.add_paragraph()

    add_line(doc, "To:")
    for line in RECIPIENT_LINES:
        add_line(doc, line)
    doc.add_paragraph()

    add_line(doc, "Respected Sir,")
    doc.add_paragraph()

    if payment_type == "Payment to Self":
        subject = f"Sub: Request for Reimbursement of Expenses for {event_name}"
        opening = (
            f"I request reimbursement for expenses I incurred for {event_name} "
            f"on {purchase_date.strftime('%d-%m-%Y')}. The expense details are:"
        )
    else:
        subject = f"Sub: Request for Vendor Payment Processing for {event_name}"
        opening = (
            f"I request direct payment to {vendor_name} for purchases made for "
            f"{event_name} on {purchase_date.strftime('%d-%m-%Y')}. The invoice details are:"
        )

    add_line(doc, subject, bold=True)
    doc.add_paragraph(opening)

    for description, amount in items:
        add_line(doc, f"•  {description} — {money(amount)}")

    doc.add_paragraph()
    add_line(doc, f"Total Amount: {money(total_amount)}")
    add_line(doc, f"Approved Budget: {money(approved_budget)}")
    doc.add_paragraph()

    if payment_type == "Payment to Self":
        doc.add_paragraph("Please process the reimbursement to the following bank account:")
        add_line(doc, f"Bank: {BANK_NAME}")
        add_line(doc, f"Account Number: {BANK_ACCOUNT_NUMBER}")
        add_line(doc, f"Account Holder: {BANK_ACCOUNT_HOLDER}")
        add_line(doc, f"Branch: {BANK_BRANCH}")
        add_line(doc, f"IFSC: {BANK_IFSC}")
    else:
        doc.add_paragraph(f"Kindly process the payment directly to {vendor_name}.")

    doc.add_paragraph("The relevant invoice(s) and supporting bills are attached for verification.")
    doc.add_paragraph("Thanking you.")
    doc.add_paragraph()
    add_line(doc, "Yours sincerely,")
    doc.add_paragraph()
    add_line(doc, SIGNATURE_LINE)
    add_line(doc, COMMITTEE)

    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output.getvalue()


st.set_page_config(page_title="SportsCom Reimbursement Generator", page_icon="📄", layout="centered")
st.title("SportsCom Reimbursement Generator")
st.caption("DoMS, IIT Madras · Generate a standardized reimbursement letter")

payment_type = st.radio(
    "Payment type",
    ["Payment to Self", "Payment to Vendor"],
    horizontal=True,
)

col1, col2 = st.columns(2)
with col1:
    event_name = st.text_input("Event / purpose", placeholder="e.g. Chak De DoMS")
with col2:
    letter_date = st.date_input("Letter date", value=date.today())

purchase_date = st.date_input("Purchase / expense date", value=date.today())

vendor_name = ""
if payment_type == "Payment to Vendor":
    vendor_name = st.text_input("Vendor name", placeholder="e.g. Asna Garments")

st.subheader("Expense items")
st.caption("Add each item once. Enter the amount for that line item.")

if "expense_rows" not in st.session_state:
    st.session_state.expense_rows = [{"description": "", "amount": ""}]

for i, row in enumerate(st.session_state.expense_rows):
    c1, c2, c3 = st.columns([5, 2, 1])
    with c1:
        row["description"] = st.text_input(
            "Description" if i == 0 else f"Description {i + 1}",
            value=row["description"],
            key=f"description_{i}",
            placeholder="e.g. Jerseys",
        )
    with c2:
        row["amount"] = st.text_input(
            "Amount (INR)" if i == 0 else f"Amount {i + 1} (INR)",
            value=row["amount"],
            key=f"amount_{i}",
            placeholder="46970",
        )
    with c3:
        st.write("")
        st.write("")
        if len(st.session_state.expense_rows) > 1 and st.button("✕", key=f"remove_{i}"):
            st.session_state.expense_rows.pop(i)
            st.rerun()

if st.button("＋ Add another item"):
    st.session_state.expense_rows.append({"description": "", "amount": ""})
    st.rerun()

col1, col2 = st.columns(2)
with col1:
    total_input = st.text_input("Total amount (INR)", placeholder="46970")
with col2:
    budget_input = st.text_input("Approved budget (INR)", placeholder="83630")

st.info("Sender, recipient, signature, committee name and (for self reimbursement) bank details are fixed in the app code.")

if st.button("Generate DOCX", type="primary", use_container_width=True):
    errors = []
    if not event_name.strip():
        errors.append("Enter the event or purpose.")
    if payment_type == "Payment to Vendor" and not vendor_name.strip():
        errors.append("Enter the vendor name.")

    items = []
    for i, row in enumerate(st.session_state.expense_rows, start=1):
        description = row["description"].strip()
        amount_text = row["amount"].strip()
        if not description and not amount_text:
            continue
        if not description or not amount_text:
            errors.append(f"Complete both description and amount for item {i}.")
            continue
        try:
            items.append((description, parse_amount(amount_text, f"item {i}")))
        except ValueError as exc:
            errors.append(str(exc))

    if not items:
        errors.append("Add at least one complete expense item.")

    try:
        total_amount = parse_amount(total_input, "total amount")
    except ValueError as exc:
        errors.append(str(exc))
        total_amount = Decimal("0")

    try:
        approved_budget = parse_amount(budget_input, "approved budget")
    except ValueError as exc:
        errors.append(str(exc))
        approved_budget = Decimal("0")

    if items and not errors:
        item_sum = sum((amount for _, amount in items), Decimal("0"))
        if item_sum != total_amount:
            errors.append(
                f"Item amounts add up to {money(item_sum)}, but Total amount is {money(total_amount)}. "
                "Please make them match."
            )

    if errors:
        for error in errors:
            st.error(error)
    else:
        file_bytes = create_letter(
            payment_type=payment_type,
            letter_date=letter_date,
            event_name=event_name.strip(),
            purchase_date=purchase_date,
            items=items,
            total_amount=total_amount,
            approved_budget=approved_budget,
            vendor_name=vendor_name.strip(),
        )
        safe_event = "".join(ch.lower() if ch.isalnum() else "_" for ch in event_name.strip()).strip("_")
        filename = f"{safe_event or 'event'}_{'self_reimbursement' if payment_type == 'Payment to Self' else 'vendor_payment'}.docx"
        st.success("Your letter is ready.")
        st.download_button(
            "Download Word document (.docx)",
            data=file_bytes,
            file_name=filename,
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            type="primary",
            use_container_width=True,
        )
